"""文档解析模块：支持 PDF / DOCX / XLSX / TXT / MD"""
import os
from typing import Tuple

import pdfplumber
from docx import Document
import openpyxl


def parse_pdf(file_path: str) -> str:
    all_text = []
    with pdfplumber.open(file_path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text()
            if text:
                all_text.append(f"[第{i+1}页]\n{text}")
    return "\n\n".join(all_text)


def parse_docx(file_path: str) -> str:
    doc = Document(file_path)
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            if para.style.name.startswith("Heading"):
                paragraphs.append(f"# {text}")
            else:
                paragraphs.append(text)
    return "\n".join(paragraphs)


def parse_xlsx(file_path: str) -> str:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    all_sheets = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            str_row = [str(cell) if cell is not None else "" for cell in row]
            if any(str_row):
                rows.append(str_row)
        if rows:
            md = f"## Sheet: {sheet_name}\n\n"
            if len(rows) > 1:
                header = "| " + " | ".join(rows[0]) + " |"
                separator = "|" + "|".join(["---" for _ in rows[0]]) + "|"
                body = "\n".join(["| " + " | ".join(r) + " |" for r in rows[1:]])
                md += f"{header}\n{separator}\n{body}"
            else:
                md += "\n".join([" | ".join(r) for r in rows])
            all_sheets.append(md)
    return "\n\n".join(all_sheets)


def parse_txt(file_path: str) -> str:
    encodings = ["utf-8", "gbk", "gb2312", "latin-1"]
    for encoding in encodings:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ValueError(f"无法识别文件编码: {file_path}")


_PARSERS = {
    ".pdf": ("PDF", parse_pdf),
    ".docx": ("DOCX", parse_docx),
    ".xlsx": ("XLSX", parse_xlsx),
    ".xls": ("XLSX", parse_xlsx),
    ".txt": ("TXT", parse_txt),
    ".md": ("TXT", parse_txt),
}


def parse_document(file_path: str) -> Tuple[str, str]:
    """根据扩展名自动选择解析器，返回 (文件类型, 文本内容)"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext not in _PARSERS:
        raise ValueError(f"不支持的文件格式: {ext}")
    doc_type, parser = _PARSERS[ext]
    return doc_type, parser(file_path)
